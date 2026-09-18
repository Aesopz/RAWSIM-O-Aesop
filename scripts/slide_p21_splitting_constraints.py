# -*- coding: utf-8 -*-
"""Slim p21: 'What makes M4G a splitting model' — three constraints, three one-line readings.
Built as native shapes on the deck template (header/footer/logos kept). Subscripts are real
PowerPoint subscript runs (markup: x_{o,i,w}).

usage: python scripts/slide_p21_splitting_constraints.py <template.pptx> <out.pptx>
"""
import sys, re
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

NAVY, NAVY_FILL = RGBColor(0x1F, 0x38, 0x64), RGBColor(0xDE, 0xEA, 0xF6)
ORANGE, ORANGE_FILL = RGBColor(0xC5, 0x5A, 0x11), RGBColor(0xFB, 0xE5, 0xD6)
GREEN, GREEN_FILL = RGBColor(0x38, 0x76, 0x1D), RGBColor(0xE2, 0xF0, 0xD9)
GREY, BLACK, WHITE = RGBColor(0x7F, 0x7F, 0x7F), RGBColor(0, 0, 0), RGBColor(0xFF, 0xFF, 0xFF)
FONT, MATH = "Times New Roman", "Cambria Math"

src, dst = sys.argv[1], sys.argv[2]
prs = Presentation(src); slide = prs.slides[0]
KEEP = {"Group 2", "Group 4", "Group 6", "TextBox 213", "TextBox 214", "TextBox 215", "TextBox 216", "TextBox 217", "TextBox 218", "TextBox 219"}
for sh in list(slide.shapes):
    if sh.name not in KEEP: sh._element.getparent().remove(sh._element)
for sh in slide.shapes:
    if sh.name == "TextBox 214":
        r = sh.text_frame.paragraphs[0].runs; r[0].text = "Proposed Model: M4G - What Makes This a Splitting Model"
        for x in r[1:]: x.text = ""
        sh.width = Emu(14500000); sh.text_frame.word_wrap = False
    if sh.name == "TextBox 219":
        r = sh.text_frame.paragraphs[0].runs
        if r: r[0].text = "第 21 頁"

W_SLIDE, H_SLIDE = prs.slide_width, prs.slide_height
def emu(inch): return int(inch * 914400)

def math_runs(p, text, size, color):
    """'Σ_{w} g_{o,i,w} = e_{o,i}' → runs with real subscripts."""
    for tok in re.split(r"(_\{[^}]*\})", text):
        if not tok: continue
        run = p.add_run()
        if tok.startswith("_{"):
            run.text = tok[2:-1]; run.font.size = Pt(size * 0.7)
            run.font._element.set("baseline", "-25000")
        else:
            run.text = tok; run.font.size = Pt(size)
        run.font.name = MATH; run.font.color.rgb = color

def textbox(x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    tb = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h)); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(60000); tf.margin_top = tf.margin_bottom = Emu(30000); tf.vertical_anchor = anchor
    first = True
    for kind, text, size, color, bold in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False; p.alignment = align
        if kind == "math": math_runs(p, text, size, color)
        else:
            r = p.add_run(); r.text = text; r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return tb

def pill(x, y, w, h, fill, edge):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, emu(x), emu(y), emu(w), emu(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill; s.line.color.rgb = edge; s.line.width = Pt(1.25); s.shadow.inherit = False
    s.adjustments[0] = 0.25; return s

# ---- one line of context under the title
textbox(0.45, 1.75, 16.0, 0.5, [("text", "Relative to M1G, one relaxation: the assignment variable moves from (order, station) to (order, line, station).  Two constraints keep splitting from being free.", 13, GREY, False)])

# ---- three rows: formula pill | reading
rows = [
    (NAVY_FILL, NAVY,   "Σ_{w∈W} g_{o,i,w} = e_{o,i}",                    "Lines are atomic",
     "A line is served in full at exactly one station. Splitting happens between the lines of an order, never inside a line."),
    (ORANGE_FILL, ORANGE, "Σ_{i} d_{o,i} g_{o,i,w} ≤ D_{o} x_{o,w}     Σ_{o} x_{o,w} ≤ C_{w}", "Splitting costs slots",
     "Each additional station an order uses consumes one of that station's scarce slots."),
    (GREEN_FILL, GREEN,  "ê_{o,i} ≥ f̂_{o}      e_{o,i} ≥ f_{o}",         "Consolidation required",
     "An order completes only when all of its lines close, wherever they landed."),
]
y0, rh, gap = 2.6, 1.85, 0.45
for k, (fill, edge, formula, head, body) in enumerate(rows):
    y = y0 + k * (rh + gap)
    pill(0.6, y, 8.2, rh, fill, edge)
    textbox(0.75, y, 7.9, rh, [("math", formula, 24, BLACK, False)], align=PP_ALIGN.CENTER)
    textbox(9.2, y, 10.4, rh, [("text", head, 20, edge, True), ("text", body, 14, BLACK, False)], anchor=MSO_ANCHOR.MIDDLE)

# ---- footer note (one line, grey)
textbox(0.45, 9.55, 16.0, 0.45, [("text", "Full constraint set (supply, robot–pod coupling, slot capacity) as in Thesis §4.2; unchanged from M1G except the relaxation above.", 11, GREY, False)])

prs.save(dst); print("->", dst)
