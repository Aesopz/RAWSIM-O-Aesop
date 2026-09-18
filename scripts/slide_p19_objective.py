# -*- coding: utf-8 -*-
"""p19 rebuilt: one chain of four questions (why D/P, why efficiency, why Dinkelbach, why D − λP)
plus one deterministic figure (F(λ) and its root). Native shapes on the deck template.

usage: python scripts/slide_p19_objective.py <template.pptx> <out.pptx>
"""
import sys, os, re, tempfile
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree

NAVY, NAVY_FILL = RGBColor(0x1F, 0x38, 0x64), RGBColor(0xDE, 0xEA, 0xF6)
ORANGE, ORANGE_FILL = RGBColor(0xC5, 0x5A, 0x11), RGBColor(0xFB, 0xE5, 0xD6)
GREY, BLACK, WHITE = RGBColor(0x7F, 0x7F, 0x7F), RGBColor(0, 0, 0), RGBColor(0xFF, 0xFF, 0xFF)
FONT, MATH = "Times New Roman", "Cambria Math"

src, dst = sys.argv[1], sys.argv[2]
prs = Presentation(src); slide = prs.slides[0]
KEEP = {"Group 2", "Group 4", "Group 6", "TextBox 213", "TextBox 214", "TextBox 215", "TextBox 216", "TextBox 217", "TextBox 218", "TextBox 219"}
for sh in list(slide.shapes):
    if sh.name not in KEEP: sh._element.getparent().remove(sh._element)
for sh in slide.shapes:
    if sh.name == "TextBox 214":
        r = sh.text_frame.paragraphs[0].runs; r[0].text = "Optimization Objective: Why a Ratio, and How It Is Solved"
        for x in r[1:]: x.text = ""
        sh.width = Emu(14500000); sh.text_frame.word_wrap = False
    if sh.name == "TextBox 219":
        r = sh.text_frame.paragraphs[0].runs
        if r: r[0].text = "第 19 頁"

def emu(v): return int(v * 914400)

def math_runs(p, text, size, color, bold=False):
    for tok in re.split(r"(_\{[^}]*\})", text):
        if not tok: continue
        run = p.add_run()
        if tok.startswith("_{"):
            run.text = tok[2:-1]; run.font.size = Pt(size * 0.7); run.font._element.set("baseline", "-25000")
        else:
            run.text = tok; run.font.size = Pt(size)
        run.font.name = MATH; run.font.color.rgb = color; run.font.bold = bold

def textbox(x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h)); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(70000); tf.margin_top = tf.margin_bottom = Emu(40000); tf.vertical_anchor = anchor
    first = True
    for kind, text, size, color, bold in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False; p.alignment = align
        if kind == "math": math_runs(p, text, size, color, bold)
        else:
            r = p.add_run(); r.text = text; r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return tb

def card(x, y, w, h, fill, edge):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, emu(x), emu(y), emu(w), emu(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill; s.line.color.rgb = edge; s.line.width = Pt(1.25); s.shadow.inherit = False
    s.adjustments[0] = 0.08; return s

def arrow_down(x, y1, y2):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(x), emu(y1), emu(x), emu(y2))
    c.line.color.rgb = NAVY; c.line.width = Pt(1.5)
    ln = c.line._get_or_add_ln(); t = etree.SubElement(ln, qn("a:tailEnd")); t.set("type", "triangle"); t.set("w", "med"); t.set("len", "med")

# ───────────────────────── left: the chain (four cards)
X, W = 0.6, 8.9
steps = [
    (NAVY_FILL, NAVY, "Why a ratio?",
     [("text", "Every dispatch spends travel (D, metres) and earns fulfilment (P, lines closed; an order counts κ lines). Efficiency is cost per unit earned.", 14, BLACK, False),
      ("math", "min  D / P        metres per line   ≡   lines per metre", 18, NAVY, True)]),
    (NAVY_FILL, NAVY, "Why not a weighted sum?",
     [("math", "min  D − w·P   requires a series of tuning experiments to set w (a metre per line),", 15, BLACK, False),
      ("text", "and the tuned value only holds for the fleet size and station load it was tuned on.", 14, BLACK, False)]),
    (ORANGE_FILL, ORANGE, "Why Dinkelbach?",
     [("text", "D / P over binary variables is nonlinear; no MILP solver takes it.", 14, BLACK, False),
      ("math", "F(λ) = min  D − λ·P        F(λ*) = 0   ⇔   λ* = min D / P", 18, ORANGE, True)]),
    (ORANGE_FILL, ORANGE, "Why it works",
     [("math", "λ_{k+1} = D(x_{k}) / P(x_{k})     the ratio the last plan achieved", 17, BLACK, False),
      ("text", "λ falls monotonically to λ*; at λ* the best plan breaks even. λ is an output with units: metres per line.", 14, BLACK, False)]),
]
y, h, gap = 1.8, 1.6, 0.34
for k, (fill, edge, head, body) in enumerate(steps):
    card(X, y, W, h, fill, edge)
    textbox(X + 0.1, y + 0.05, W - 0.2, h - 0.1, [("text", head, 16, edge, True)] + body)
    if k < len(steps) - 1: arrow_down(X + W / 2, y + h, y + h + gap)
    y += h + gap

# ───────────────────────── right: F(λ) figure (deterministic matplotlib)
lam = np.linspace(0, 10, 400); lam_star = 4.0
# F(λ) = min over plans of D − λP: 0 (empty plan) until λ*, then the best non-empty plan's line
F = np.where(lam < lam_star, 0.0, -(lam - lam_star) * 1.0)
fig, ax = plt.subplots(figsize=(5.2, 4.2))
plt.rcParams.update({"font.family": "sans-serif", "font.size": 11})
ax.plot(lam, F, color="black", lw=1.8)
ax.axhline(0, color="#7F7F7F", lw=0.8)
ax.plot([lam_star], [0], marker="o", color="#C55A11", ms=8, zorder=5)
ax.annotate("λ*  =  min D/P\nbreak-even, best plan non-empty", (lam_star, 0), xytext=(lam_star + 0.6, 1.6), fontsize=10,
            arrowprops=dict(arrowstyle="-", color="#C55A11", lw=0.8), color="#C55A11")
ax.text(1.8, 0.55, "λ < λ*: nothing pays,\nplan is empty (P = 0)", fontsize=10, ha="center", color="#1F3864")
ax.text(5.6, -4.8, "λ > λ*: over-priced,\nF(λ) < 0, over-dispatch", fontsize=10, ha="center", color="#1F3864")
ax.set_xlabel("exchange rate λ  (metres per line)"); ax.set_ylabel("F(λ) = min D − λ·P")
ax.set_xticks([lam_star]); ax.set_xticklabels(["λ*"]); ax.set_yticks([0]); ax.set_ylim(-6.5, 2.6); ax.set_xlim(0, 10)
for s in ("top", "right"): ax.spines[s].set_visible(False)
png = os.path.join(tempfile.gettempdir(), "p19_F_lambda.png"); fig.tight_layout(); fig.savefig(png, dpi=300); plt.close(fig)
slide.shapes.add_picture(png, emu(10.3), emu(1.8), width=emu(8.4))
textbox(10.3, 8.65, 8.4, 0.8, [("text", "Root-finding is Newton on F: 2–3 solves per decision.  The exchange rate is measured by the model, not chosen by the modeller.", 12, GREY, False)])

prs.save(dst); print("->", dst)
