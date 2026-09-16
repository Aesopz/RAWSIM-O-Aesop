# -*- coding: utf-8 -*-
"""Exact statistics pipeline (docs/REPORTING-STANDARD.md section 5).

Every number reported to the user is computed here, never by hand:
  1. per-seed metrics read from raw engine output (full float64 precision, no rounding)
  2. descriptives and paired tests with numpy/scipy (exact t distribution p-values, CIs)
  3. an Excel workbook whose check sheet recomputes means, SDs, Δ% and paired T.TEST with
     native Excel formulas; Excel is launched headless to recalculate, the values are read back
     and compared with scipy (relative tolerance 1e-9). Any mismatch aborts.
  4. APA 7 outputs: stats.xlsx, stats.json, apa_tables.docx (three-line tables, English)
Rounding happens only at presentation, half-up (Decimal), never before a computation.

Usage
  python scripts/stats_pipeline.py exp <experiment_id>            (called by exp.py table)
  python scripts/stats_pipeline.py legacy <out_dir> <spec.json>   (pre-exp.py data under out/)
     spec.json = {"title": "...", "scenario_note": "...",
                  "arms": [{"label": "M1G", "bots": 6, "runs": {"0": "p5_m1g_6_s0", ...}}],
                  "comparisons": [{"a": "M1G", "b": "M4G-NS"}]}
"""
import csv, datetime, glob, io, json, math, os, re, subprocess, sys
from decimal import Decimal, ROUND_HALF_UP

import numpy as np
from scipy import stats as sps

ROOT = r"C:\Users\Aesop\Desktop\EE-RAWSim-O_PP"
METRICS = ["Items", "Lines", "Orders", "Pile-on", "Trips", "Trips/Orders", "m/Line",
           "EOR(kJ/order)", "TurnoverMedian(s)", "StationIdle(%)"]
APA_HEAD = {"Items": "Items", "Lines": "Lines", "Orders": "Orders", "Pile-on": "Pile-on", "Trips": "Trips",
            "Trips/Orders": "Trips/order", "m/Line": "m/line", "EOR(kJ/order)": "EOR (kJ/order)",
            "TurnoverMedian(s)": "Turnover (s)", "StationIdle(%)": "Station idle (%)"}
DECIMALS = {"Items": 1, "Lines": 1, "Orders": 1, "Pile-on": 3, "Trips": 1, "Trips/Orders": 4,
            "m/Line": 2, "EOR(kJ/order)": 3, "TurnoverMedian(s)": 1, "StationIdle(%)": 2}
REL_TOL = 1e-9
EXCEL = r"C:\Program Files\Microsoft Office\root\Office16\EXCEL.EXE"


# ---------------------------------------------------------------- metrics from raw output
def _read(path):
    return io.open(path, encoding="utf-8-sig", errors="ignore").read()

def run_metrics(run_dir):
    """run_dir contains exactly one engine output folder with statistics.txt."""
    hits = glob.glob(os.path.join(run_dir, "*", "statistics.txt"))
    if len(hits) != 1:
        raise SystemExit("expected one statistics.txt under %s, found %d" % (run_dir, len(hits)))
    txt = _read(hits[0])
    def g(key):
        m = re.search(r"^" + re.escape(key) + r": ([-\d.E+]+)\s*$", txt, re.M)
        if not m:
            raise SystemExit("%s missing in %s" % (key, hits[0]))
        return float(m.group(1))
    sub = os.path.dirname(hits[0])
    items, lines, orders = g("StatOverallItemsHandled"), g("StatOverallLinesHandled"), g("StatOverallOrdersHandled")
    trips = g("StatOutputStationArrivals")
    idle = [float(r["IdleTime"]) / float(r["UpTime"]) * 100.0
            for r in csv.DictReader(io.open(os.path.join(sub, "stationstatistics.csv"), encoding="utf-8"), delimiter=";")
            if r["Ident"].startswith("OutputStation")]
    return {"Items": items, "Lines": lines, "Orders": orders,
            "Pile-on": orders / trips, "Trips": trips, "Trips/Orders": trips / orders,
            "m/Line": g("StatOverallDistanceTraveled") / lines, "EOR(kJ/order)": g("KPI_EOR"),
            "TurnoverMedian(s)": g("StatMedianTurnoverTime"), "StationIdle(%)": float(np.mean(idle)),
            "_hours": orders / g("StatThroughputOrdersPerHour"), "_source": os.path.relpath(hits[0], ROOT)}


# ---------------------------------------------------------------- statistics
def describe(values):
    x = np.asarray(values, dtype=np.float64)
    n = x.size
    mean = float(x.mean())
    sd = float(x.std(ddof=1)) if n > 1 else float("nan")
    se = sd / math.sqrt(n) if n > 1 else float("nan")
    half = float(sps.t.ppf(0.975, n - 1)) * se if n > 1 else float("nan")
    return {"n": n, "mean": mean, "sd": sd, "se": se, "ci_low": mean - half, "ci_high": mean + half,
            "min": float(x.min()), "max": float(x.max())}

def paired(a_vals, b_vals):
    """delta = b - a, paired by seed. pct = (mean_b - mean_a) / mean_a * 100."""
    a = np.asarray(a_vals, dtype=np.float64); b = np.asarray(b_vals, dtype=np.float64)
    d = b - a
    n = d.size
    out = {"n": n, "mean_a": float(a.mean()), "mean_b": float(b.mean()), "mean_diff": float(d.mean()),
           "pct": float((b.mean() - a.mean()) / a.mean() * 100.0) if a.mean() != 0 else float("nan")}
    if n < 2:
        out.update(sd_diff=float("nan"), t=float("nan"), df=0, p=float("nan"), dz=float("nan"),
                   ci_low=float("nan"), ci_high=float("nan"), label="1 seed")
        return out
    sd = float(d.std(ddof=1))
    out["sd_diff"], out["df"] = sd, n - 1
    if sd == 0.0:
        out.update(t=float("nan"), p=float("nan"), dz=float("nan"), ci_low=out["mean_diff"], ci_high=out["mean_diff"],
                   label="identical" if out["mean_diff"] == 0 else "constant difference")
        return out
    res = sps.ttest_rel(b, a)
    half = float(sps.t.ppf(0.975, n - 1)) * sd / math.sqrt(n)
    p = float(res.pvalue)
    out.update(t=float(res.statistic), p=p, dz=out["mean_diff"] / sd,
               ci_low=out["mean_diff"] - half, ci_high=out["mean_diff"] + half,
               label="***" if p < .001 else "**" if p < .01 else "*" if p < .05 else "")
    return out


# ---------------------------------------------------------------- presentation rounding (half-up)
def fmt(x, dec):
    if x is None or (isinstance(x, float) and (math.isnan(x) or math.isinf(x))):
        return ""
    q = Decimal(1).scaleb(-dec) if dec > 0 else Decimal(1)
    s = str(Decimal(repr(float(x))).quantize(q, rounding=ROUND_HALF_UP))
    return s.replace("-", "\u2212")

def fmt_signed(x, dec):
    s = fmt(x, dec)
    return s if (not s or s.startswith("\u2212") or float(x) == 0) else "+" + s

def fmt_p(p):
    if p is None or math.isnan(p):
        return ""
    return "< .001" if p < .001 else fmt(p, 3).lstrip("0")


# ---------------------------------------------------------------- Excel cross-check
def write_workbook(path, arms, comparisons, results):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter as L
    wb = Workbook()
    ws = wb.active; ws.title = "per_seed"
    ws.append(["arm", "bots", "seed"] + METRICS)
    ranges = {}                                          # (arm, metric) -> "per_seed!$D$2:$D$11"
    row = 2
    for arm in arms:
        first = row
        for seed in arm["seeds"]:
            m = arm["data"][seed]
            ws.append([arm["label"], arm["bots"], seed] + [m[k] for k in METRICS])
            row += 1
        for j, k in enumerate(METRICS):
            col = L(4 + j)
            ranges[(arm["label"], k)] = "per_seed!$%s$%d:$%s$%d" % (col, first, col, row - 1)
    ck = wb.create_sheet("excel_check")
    ck.append(["kind", "arm_or_comparison", "metric", "excel_value", "python_value"])
    checks = []
    for arm in arms:
        for k in METRICS:
            r = ranges[(arm["label"], k)]
            d = results["descriptives"][arm["label"]][k]
            checks.append(("mean", arm["label"], k, "=AVERAGE(%s)" % r, d["mean"]))
            if d["n"] > 1:
                checks.append(("sd", arm["label"], k, "=_xlfn.STDEV.S(%s)" % r, d["sd"]))
    for c in comparisons:
        name = "%s vs %s" % (c["b"], c["a"])
        for k in METRICS:
            ra, rb = ranges[(c["a"], k)], ranges[(c["b"], k)]
            pr = results["paired"][name][k]
            checks.append(("pct", name, k, "=(AVERAGE(%s)-AVERAGE(%s))/AVERAGE(%s)*100" % (rb, ra, ra), pr["pct"]))
            if pr["n"] > 1 and not math.isnan(pr["p"]):
                checks.append(("p", name, k, "=_xlfn.T.TEST(%s,%s,2,1)" % (ra, rb), pr["p"]))
    for chk in checks:
        ck.append(list(chk))
    for sheet in (ws, ck):
        for i in range(1, sheet.max_column + 1):
            sheet.column_dimensions[L(i)].width = 18
    wb.save(path)
    return len(checks)

def excel_recalculate(path):
    if not os.path.exists(EXCEL):
        return False, "Excel not installed"
    ps = ("$ErrorActionPreference='Stop'; $x=New-Object -ComObject Excel.Application; $x.Visible=$false; "
          "$x.DisplayAlerts=$false; try { $wb=$x.Workbooks.Open('%s'); $x.CalculateFull(); $wb.Save(); $wb.Close($true) } "
          "finally { $x.Quit(); [System.Runtime.InteropServices.Marshal]::ReleaseComObject($x) | Out-Null }" % path.replace("'", "''"))
    r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", ps], capture_output=True, text=True, timeout=300)
    return r.returncode == 0, (r.stderr or "").strip()

def excel_compare(path):
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True)["excel_check"]
    bad, n = [], 0
    for kind, name, metric, xv, pv in ws.iter_rows(min_row=2, values_only=True):
        n += 1
        if not isinstance(xv, (int, float)):
            bad.append("%s %s %s: excel=%r" % (kind, name, metric, xv)); continue
        if not math.isclose(float(xv), float(pv), rel_tol=REL_TOL, abs_tol=1e-12):
            bad.append("%s %s %s: excel=%.15g python=%.15g" % (kind, name, metric, xv, pv))
    return n, bad


# ---------------------------------------------------------------- APA 7 docx tables
def _three_line(table):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    def border(cell, **edges):
        tcPr = cell._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            el = OxmlElement("w:%s" % edge)
            if edge in edges:
                el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(edges[edge])); el.set(qn("w:color"), "000000")
            else:
                el.set(qn("w:val"), "nil")
            b.append(el)
        tcPr.append(b)
    last = len(table.rows) - 1
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            e = {}
            if i == 0: e.update(top=8, bottom=4)
            if i == last: e["bottom"] = 8
            border(cell, **e)

def _add_table(doc, number, title, header, rows, note):
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    p = doc.add_paragraph(); r = p.add_run("Table %d" % number); r.bold = True
    p = doc.add_paragraph(); r = p.add_run(title); r.italic = True
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    for j, h in enumerate(header):
        t.rows[0].cells[j].text = h
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            t.rows[i].cells[j].text = v
            if j >= 2:
                t.rows[i].cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for row in t.rows:
        for cell in row.cells:
            for par in cell.paragraphs:
                for run in par.runs:
                    run.font.size = Pt(9); run.font.name = "Times New Roman"
    _three_line(t)
    if note:
        p = doc.add_paragraph(); r = p.add_run("Note. "); r.italic = True; p.add_run(note)
    doc.add_paragraph()

def write_apa_docx(path, title, scenario_note, arms, comparisons, results):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    header = ["Policy", "Bots"] + [APA_HEAD[k] for k in METRICS]
    rows = [[a["label"], str(a["bots"])] + [fmt(results["descriptives"][a["label"]][k]["mean"], DECIMALS[k]) for k in METRICS]
            for a in arms]
    _add_table(doc, 1, title, header, rows, scenario_note + " Values are means across seeds.")
    sd_rows = [[a["label"], str(a["bots"])] + [fmt(results["descriptives"][a["label"]][k]["sd"], DECIMALS[k]) for k in METRICS]
               for a in arms]
    if all(a["n"] > 1 for a in arms):
        _add_table(doc, 2, title + " (Standard Deviations)", header, sd_rows, None)
    if comparisons:
        n = arms[0]["n"]
        rows = []
        for c in comparisons:
            name = "%s vs %s" % (c["b"], c["a"])
            cells = []
            for k in METRICS:
                pr = results["paired"][name][k]
                star = pr["label"] if pr["label"] in ("*", "**", "***") else ""
                cells.append(fmt_signed(pr["pct"], 2) + star)
            rows.append([c["b"], c["a"]] + cells)
        note = ("Percentage difference of the first policy relative to the second, paired by seed (n = %d)." % n
                if n > 1 else "Percentage difference from a single seed; no significance test.")
        if n > 1:
            note += " Two-tailed paired t tests. *p < .05. **p < .01. ***p < .001."
        _add_table(doc, 3, "Paired Differences (%)", ["Policy", "Reference"] + [APA_HEAD[k] for k in METRICS], rows, note)
    doc.save(path)


# ---------------------------------------------------------------- pipeline
def run_pipeline(out_dir, title, scenario_note, arms, comparisons):
    """arms: [{"label", "bots", "seeds": [..], "dirs": {seed: run_dir}}]"""
    os.makedirs(out_dir, exist_ok=True)
    for a in arms:
        a["seeds"] = sorted(a["seeds"])                  # Excel ranges and scipy pairing share this order
        a["data"] = {s: run_metrics(a["dirs"][s]) for s in a["seeds"]}
        a["n"] = len(a["seeds"])
    for c in comparisons:
        sa = next(x for x in arms if x["label"] == c["a"])["seeds"]
        sb = next(x for x in arms if x["label"] == c["b"])["seeds"]
        if sorted(sa) != sorted(sb):
            raise SystemExit("comparison %s vs %s: seed sets differ" % (c["b"], c["a"]))
    results = {"generated": datetime.datetime.now().isoformat(timespec="seconds"),
               "definitions": {"pct": "(mean_b - mean_a) / mean_a * 100", "delta": "b - a paired by seed",
                               "test": "scipy.stats.ttest_rel, two-tailed", "ci": "95% t interval",
                               "dz": "mean(delta) / sd(delta)", "rounding": "presentation only, half-up"},
               "descriptives": {a["label"]: {k: describe([a["data"][s][k] for s in a["seeds"]]) for k in METRICS} for a in arms},
               "paired": {}, "per_seed": {a["label"]: {str(s): a["data"][s] for s in a["seeds"]} for a in arms}}
    for c in comparisons:
        A = next(x for x in arms if x["label"] == c["a"]); B = next(x for x in arms if x["label"] == c["b"])
        seeds = sorted(A["seeds"])
        results["paired"]["%s vs %s" % (c["b"], c["a"])] = {
            k: paired([A["data"][s][k] for s in seeds], [B["data"][s][k] for s in seeds]) for k in METRICS}
    xlsx = os.path.join(out_dir, "stats.xlsx")
    n_checks = write_workbook(xlsx, arms, comparisons, results)
    ok, err = excel_recalculate(xlsx)
    if ok:
        n, bad = excel_compare(xlsx)
        if bad:
            raise SystemExit("EXCEL CROSS-CHECK FAILED (%d/%d):\n  %s" % (len(bad), n, "\n  ".join(bad[:20])))
        results["excel_check"] = {"status": "passed", "checks": n, "rel_tol": REL_TOL}
    else:
        results["excel_check"] = {"status": "not run", "reason": err[:300], "checks_written": n_checks}
    with io.open(os.path.join(out_dir, "stats.json"), "w", encoding="utf-8") as f:
        json.dump(results, f, indent=1, ensure_ascii=False, default=float)
    write_apa_docx(os.path.join(out_dir, "apa_tables.docx"), title, scenario_note, arms, comparisons, results)
    return results

def summary_text(arms, comparisons, results, scenario_line):
    """Plain summary for results.csv / NOTES.md (precise values, presentation rounding)."""
    lines = [scenario_line, "policy,bots," + ",".join(METRICS)]
    for a in arms:
        lines.append("%s,%db,%s" % (a["label"], a["bots"], ",".join(
            fmt(results["descriptives"][a["label"]][k]["mean"], DECIMALS[k]) for k in METRICS)))
    for c in comparisons:
        name = "%s vs %s" % (c["b"], c["a"])
        lines += ["", "paired %s (delta%% = (mean_b - mean_a)/mean_a; n=%d)" % (name, arms[0]["n"]),
                  "stat," + ",".join(METRICS)]
        P = results["paired"][name]
        lines.append("pct," + ",".join(fmt_signed(P[k]["pct"], 2) for k in METRICS))
        lines.append("p," + ",".join(fmt_p(P[k]["p"]) or P[k]["label"] for k in METRICS))
        lines.append("sig," + ",".join(P[k]["label"] for k in METRICS))
    lines += ["", "excel cross-check: %s" % json.dumps(results["excel_check"], ensure_ascii=False)]
    return "\n".join(lines)


def main():
    if len(sys.argv) < 3:
        print(__doc__); sys.exit(1)
    if sys.argv[1] == "exp":
        sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
        import exp as E
        E.cmd_table(sys.argv[2])
    elif sys.argv[1] == "legacy":
        out_dir = os.path.abspath(sys.argv[2])
        spec = json.load(io.open(sys.argv[3], encoding="utf-8"))
        arms = [{"label": a["label"], "bots": a["bots"], "seeds": sorted(int(s) for s in a["runs"]),
                 "dirs": {int(s): os.path.join(ROOT, d) if d.replace("\\", "/").split("/")[0] in ("out", "experiments") else os.path.join(ROOT, "out", d) for s, d in a["runs"].items()}} for a in spec["arms"]]
        res = run_pipeline(out_dir, spec["title"], spec["scenario_note"], arms, spec.get("comparisons", []))
        print(summary_text(arms, spec.get("comparisons", []), res, spec["scenario_note"]))
    else:
        print(__doc__); sys.exit(1)

if __name__ == "__main__":
    main()
