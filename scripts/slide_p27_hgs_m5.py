# -*- coding: utf-8 -*-
"""Rebuild slide p27 (HGS-M5 flowchart) as NATIVE PowerPoint shapes in the deck's own style.

Keeps the template's header/footer/logos, deletes the old body, and draws the Canon v2 inner loop
(draws first, dispatch only when no draw is negative, whole-order fallback when the split budget
binds). Same geometry as scripts/fig_m5_flowchart.py, so the PNG and the PPTX agree.

usage: python scripts/slide_p27_hgs_m5.py <template.pptx> <out.pptx>
"""
import sys, copy
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree

NAVY, NAVY_FILL = RGBColor(0x1F, 0x38, 0x64), RGBColor(0xDE, 0xEA, 0xF6)
ORANGE, ORANGE_FILL = RGBColor(0xC5, 0x5A, 0x11), RGBColor(0xFB, 0xE5, 0xD6)
GREY, GREY_FILL, BLACK, WHITE = RGBColor(0x7F, 0x7F, 0x7F), RGBColor(0xF2, 0xF2, 0xF2), RGBColor(0, 0, 0), RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Times New Roman"

src, dst = sys.argv[1], sys.argv[2]
prs = Presentation(src)
slide = prs.slides[0]

# ---- keep header / footer / logos, drop the old body
KEEP = {"Group 2", "Group 4", "Group 6", "TextBox 213", "TextBox 214", "TextBox 215", "TextBox 216", "TextBox 217", "TextBox 218", "TextBox 219"}
for sh in list(slide.shapes):
    if sh.name not in KEEP:
        sh._element.getparent().remove(sh._element)
for sh in slide.shapes:
    if sh.name == "TextBox 214":
        r = sh.text_frame.paragraphs[0].runs
        r[0].text = "Proposed Heuristic Model : HGS-M5"
        for extra in r[1:]: extra.text = ""

# ---- coordinate mapping: figure units (0..133.3 x, 0..62 y, y up) -> EMU
S = 131000; OX = 350000; OY = 1200000
def X(u): return int(OX + u * S)
def Y(u): return int(OY + (62 - u) * S)
def W(u): return int(u * S)

def _style_text(tf, lines, size, color, bold, align=PP_ALIGN.CENTER):
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(40000); tf.margin_top = tf.margin_bottom = Emu(20000)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    first = True
    for text, sz, col, b in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        run = p.add_run(); run.text = text
        run.font.name = FONT; run.font.size = Pt(sz); run.font.bold = b; run.font.color.rgb = col

def box(x, y, w, h, title, sub=None, fill=NAVY_FILL, edge=NAVY, bold=True, fs=13, sfs=10.5, tcolor=BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = slide.shapes.add_shape(shape, X(x - w / 2), Y(y + h / 2), W(w), W(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = edge; sh.line.width = Pt(1.25)
    sh.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = 0.12
    lines = [(title, fs, tcolor, bold)] + ([(sub, sfs, tcolor, False)] if sub else [])
    _style_text(sh.text_frame, lines, fs, tcolor, bold)
    return sh

def diamond(x, y, w, h, text, fs=12):
    sh = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, X(x - w / 2), Y(y + h / 2), W(w), W(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = WHITE
    sh.line.color.rgb = ORANGE; sh.line.width = Pt(1.5); sh.shadow.inherit = False
    _style_text(sh.text_frame, [(text, fs, NAVY, True)], fs, NAVY, True)
    return sh

def label(x, y, text, color=ORANGE, fs=10.5, ha="center", bold=False, w=None):
    lines = text.split("\n")
    est_w = w if w else max(len(l) for l in lines) * fs * 0.55 / 72 * 914400 + 100000   # EMU estimate
    h = int(len(lines) * fs * 1.25 / 72 * 914400 + 60000)
    left = X(x) - (est_w // 2 if ha == "center" else (0 if ha == "left" else est_w))
    tb = slide.shapes.add_textbox(int(left), Y(y) - h // 2, int(est_w), h)
    tf = tb.text_frame; tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    align = {"center": PP_ALIGN.CENTER, "left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT}[ha]
    _style_text(tf, [(l, fs, color, bold) for l in lines], fs, color, bold, align)
    return tb

def _arrowhead(conn):
    ln = conn.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd")); tail.set("type", "triangle"); tail.set("w", "med"); tail.set("len", "med")

def seg(p, q, color=NAVY, head=False, lw=1.25):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, X(p[0]), Y(p[1]), X(q[0]), Y(q[1]))
    c.line.color.rgb = color; c.line.width = Pt(lw)
    if head: _arrowhead(c)
    return c

def elbow(pts, color=NAVY, head=True):
    for i in range(len(pts) - 1):
        seg(pts[i], pts[i + 1], color, head=(head and i == len(pts) - 2))

def arrow(p, q, color=NAVY): seg(p, q, color, head=True)

# ───────────────────────── column headers
label(3, 59.5, "OUTER  Updating λ", GREY, 12, ha="left")
label(50, 59.5, "INNER  BuildPlan(λ)", NAVY, 12, ha="left", bold=True)
label(103, 59.5, "Move Scores", GREY, 12, ha="left")
frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X(48.5), Y(57.7), W(50), W(54.2))
frame.fill.background(); frame.line.color.rgb = GREY; frame.line.width = Pt(1); frame.line.dash_style = 4; frame.adjustments[0] = 0.02; frame.shadow.inherit = False

# ───────────────────────── OUTER loop (unchanged from M4G)
Xo = 20; Wo = 33
box(Xo, 56, Wo, 3.6, "λ ← λ₀", fill=ORANGE_FILL, edge=ORANGE, bold=False)
box(Xo, 50.5, Wo, 4.4, "plan ← BuildPlan(λ)", "the inner loop")
box(Xo, 44.5, Wo, 3.6, "P = (D* − objective) / λ", bold=False)
diamond(Xo, 38.5, 22, 5.0, "P ≤ 0 ?")
diamond(Xo, 31.5, 22, 5.0, "| objective | ≤ tol ?")
box(Xo, 25, Wo, 3.6, "λ_next = D* / P", bold=False)
diamond(Xo, 19, 22, 5.0, "λ_next ≤ 0 ?")
box(Xo, 12.8, Wo, 3.6, "trial ← BuildPlan(λ_next)", bold=False)
diamond(Xo, 6.8, 22, 5.0, "trial P ≤ 0 ?")
for a, b in ((54.2, 52.7), (48.3, 46.3), (42.7, 41.0), (36.0, 34.0), (29.0, 26.8), (23.2, 21.5), (16.5, 14.6), (11.0, 9.3)):
    arrow((Xo, a), (Xo, b))
for c in (38.5, 31.5, 19.0, 6.8):
    label(Xo + 1.0, c - 3.1, "no", GREY, ha="left")
diamond(43, 38.5, 12, 4.4, "λ̄ > λ ?", fs=11)
arrow((31, 38.5), (37, 38.5), ORANGE); label(33.5, 39.4, "yes")
elbow([(43, 40.7), (43, 50.5), (36.6, 50.5)], ORANGE); label(44.2, 45.5, "yes\nλ ← λ̄", ha="left", fs=10)
label(43.8, 35.6, "no", GREY, ha="left"); label(46.5, 33.3, "Stop: no feasible\nnon-empty plan", ORANGE, 9.5, ha="right", bold=True)
label(32.0, 32.3, "yes", ha="left"); label(35.0, 29.6, "Stop: break-even", ORANGE, 10.5, ha="left", bold=True)
label(32.0, 19.8, "yes", ha="left"); label(35.0, 17.2, "Stop: D* = 0, no dispatch", ORANGE, 10.5, ha="left", bold=True)
label(32.0, 7.6, "yes", ha="left"); label(35.0, 5.0, "Stop: keep previous plan", ORANGE, 10.5, ha="left", bold=True)
elbow([(Xo, 4.3), (Xo, 2.2), (1.8, 2.2), (1.8, 50.5), (3.5, 50.5)])
label(2.8, 0.9, "No → adopt λ ← λ_next, re-run BuildPlan (cap 5)", GREY, 10, ha="left")

# ───────────────────────── INNER BuildPlan (Canon v2)
IX = 73.5; IW = 40
box(IX, 55.2, IW, 3.4, "Copy the working books")
box(IX, 49.6, IW, 4.6, "①  Enumerate draw-line moves", "Station stock covers the line IN FULL · slot available")
diamond(IX, 43.3, 24, 4.6, "any draw Δ < 0 ?")
box(IX, 37.3, IW, 4.6, "②  Score every dispatch candidate", "Marginal: harvest WITH the pod − WITHOUT it", fill=ORANGE_FILL, edge=ORANGE)
diamond(IX, 31.0, 26, 4.6, "best dispatch Δ < 0 ?", fs=11.5)
diamond(IX, 24.4, 24, 4.6, "split budget bound ?", fs=11)
box(IX, 18.2, IW, 4.6, "③  Whole-order fallback", "Cheapest pod set covering ONE unsplit order · any cover accepted", fill=ORANGE_FILL, edge=ORANGE)
box(IX, 11.4, IW, 4.6, "Valuation sweep", "Coverable but slot-blocked:  −λβ per line,  −μβ per order")
box(IX, 5.6, IW, 3.4, "return  plan  (D*, objective)", fill=GREY_FILL, edge=GREY, bold=False)
arrow((IX, 53.5), (IX, 51.9)); arrow((IX, 47.3), (IX, 45.6))
arrow((IX, 41.0), (IX, 39.6)); label(IX + 0.9, 40.4, "no", GREY, ha="left")
arrow((IX, 35.0), (IX, 33.3))
arrow((IX, 28.7), (IX, 26.7)); label(IX + 0.9, 27.8, "no", GREY, ha="left")
arrow((IX, 22.1), (IX, 20.5)); label(IX + 0.9, 21.4, "yes", ORANGE, ha="left")
NX = IX + 23.5
elbow([(IX + 12, 24.4), (NX, 24.4), (NX, 11.4), (IX + IW / 2, 11.4)]); label(NX - 0.6, 19.0, "no", GREY, ha="right")
arrow((IX, 9.1), (IX, 7.3))
RX = IX + 22.0
elbow([(IX + 12, 43.3), (RX, 43.3), (RX, 49.6), (IX + IW / 2, 49.6)], ORANGE)
label(IX + 12.4, 44.3, "yes: best draw", ORANGE, 10, ha="left")
elbow([(IX + 13, 31.0), (RX, 31.0), (RX, 43.3)], ORANGE, head=False)
label(IX + 13.4, 32.0, "yes: dispatch", ORANGE, 10, ha="left")
elbow([(IX - IW / 2, 18.2), (IX - 22.5, 18.2), (IX - 22.5, 49.6), (IX - IW / 2, 49.6)], ORANGE)
label(IX - 22.0, 34.5, "bundle accepted", ORANGE, 10, ha="left")
label(RX - 0.6, 46.6, "books changed", GREY, 10, ha="right")

# ───────────────────────── Move scores (right column)
SX = 115; SW = 32
def score_card(yc, h, edge, title, rows, foot=None):
    box(SX, yc, SW, h, "", fill=WHITE, edge=edge)
    top = yc + h / 2 - 1.9
    label(SX - SW / 2 + 1.2, top, title, NAVY, 11.5, ha="left", bold=True)
    yy = top - 3.0
    for formula, note in rows:
        label(SX - SW / 2 + 1.2, yy, formula, NAVY, 12.5, ha="left")
        label(SX + SW / 2 - 1.2, yy, note, ORANGE, 10, ha="right")
        yy -= 2.7
    if foot: label(SX - SW / 2 + 1.2, yy + 0.4, foot, GREY, 10, ha="left")
score_card(51.5, 9.0, NAVY, "DRAW from the committed pods in station",
           [("Δ = −λ(1 − β)", "Close an order-line"), ("Δ = −(λ + μ)(1 − β)", "Close an order")])
score_card(38.8, 12.4, ORANGE, "DISPATCH a robot to fetch the new pod",
           [("Δ = (d_r,p + d_p,w)", "Extract travel distance"), ("     + (U₊ − U₀)", "Marginal unlock"), ("     + (V₊ − V₀)", "Valuation unlock")])
score_card(25.4, 11.4, ORANGE, "WHOLE-ORDER fallback (budget bound)",
           [("Δ = Σ (d_r,p + d_p,w)", "One bot per pod"), ("accept any cover, cheapest wins", "")],
           foot="no box consumed · next decision re-enters ①")

prs.save(dst)
print("->", dst)
