"""Editable PowerPoint flowchart of the lambda update loop.

Every element is a native PowerPoint shape or connector, so the deck can be
edited normally. No rendered images, no analysis figures - process only.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

W = os.path.dirname(os.path.abspath(__file__))

NAVY = RGBColor(0x1E, 0x27, 0x61)
GOLD = RGBColor(0xB8, 0x80, 0x1F)
WARN = RGBColor(0x9C, 0x4F, 0x26)
GREY = RGBColor(0x5B, 0x63, 0x7D)
LINE = RGBColor(0x44, 0x4C, 0x66)
PAPER = RGBColor(0xF2, 0xF4, 0xFA)
GOLDBG = RGBColor(0xFB, 0xF3, 0xE2)
WARNBG = RGBColor(0xF9, 0xF0, 0xE9)
PANEL = RGBColor(0xE6, 0xEA, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

HEAD, BODY = "Cambria", "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid()
s.background.fill.fore_color.rgb = WHITE


def text_in(shape, lines, size=11, color=NAVY, bold_first=False, font=BODY):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(1)
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size if i == 0 else size - 1.5)
        r.font.bold = bold_first and i == 0
        r.font.color.rgb = color if i == 0 else GREY
        r.font.name = font


def box(x, y, w, h, lines, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        fill=WHITE, edge=LINE, size=11, color=NAVY, bold=False, width=1.0):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = edge; sh.line.width = Pt(width)
    sh.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = 0.12
    text_in(sh, lines, size=size, color=color, bold_first=bold)
    return sh


def arrow(x1, y1, x2, y2, color=LINE, dash=False, width=1.0):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                               Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn('a:tailEnd'),
                             {'type': 'triangle', 'w': 'med', 'len': 'med'}))
    if dash:
        ln.append(ln.makeelement(qn('a:prstDash'), {'val': 'dash'}))
    return c


def elbow(pts, color=LINE, dash=True, width=1.0):
    """Polyline drawn as a chain of connectors; arrowhead on the last leg only."""
    for i in range(len(pts) - 1):
        (x1, y1), (x2, y2) = pts[i], pts[i + 1]
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                   Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = color
        c.line.width = Pt(width)
        ln = c.line._get_or_add_ln()
        if dash:
            ln.append(ln.makeelement(qn('a:prstDash'), {'val': 'dash'}))
        if i == len(pts) - 2:
            ln.append(ln.makeelement(qn('a:tailEnd'),
                                     {'type': 'triangle', 'w': 'med', 'len': 'med'}))


def label(x, y, w, txt, size=9.5, color=GREY, align=PP_ALIGN.CENTER, italic=False):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.24))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = txt
    r.font.size = Pt(size); r.font.color.rgb = color
    r.font.name = BODY; r.font.italic = italic


# ── title ──────────────────────────────────────────────────────────────
label(0.6, 0.30, 8.0, "Chapter 05 · Methodology", size=10, color=GOLD,
      align=PP_ALIGN.LEFT)
tb = s.shapes.add_textbox(Inches(0.6), Inches(0.52), Inches(9.5), Inches(0.5))
p = tb.text_frame.paragraphs[0]
r = p.add_run(); r.text = "Updating λ within one decision epoch"
r.font.size = Pt(26); r.font.bold = True; r.font.color.rgb = NAVY; r.font.name = HEAD

# ── geometry ───────────────────────────────────────────────────────────
CX, BW, BH = 6.35, 3.10, 0.52          # main column
DX, DW, DH = 6.35, 3.10, 0.72          # diamonds
EX, EW, EH = 10.20, 3.30, 0.60         # exits on the right
L = CX - BW / 2.0                       # left edge of main column
EL = EX - EW / 2.0

rows = [1.24, 1.99, 2.81, 3.69, 4.49, 5.27, 6.09]

# 1 start
box(L, rows[0], BW, BH,
    ["λ ← λ₀", "measured from cumulative history"],
    fill=GOLDBG, edge=GOLD, color=GOLD, bold=True, width=1.4)

# 2 solve + recover
arrow(CX, rows[0] + BH, CX, rows[1] - 0.02)
box(L, rows[1], BW, 0.62,
    ["Solve  min D − V  at λ",
     "obtain D*, objective;   V* = (D* − objective) / λ"],
    fill=PAPER, edge=LINE)

# 3 diamond: V* <= 0
arrow(CX, rows[1] + 0.62, CX, rows[2] - 0.02)
box(L, rows[2], DW, DH, ["V*  ≤  0 ?"], shape=MSO_SHAPE.DIAMOND, size=11.5)
arrow(CX + DW / 2.0, rows[2] + DH / 2.0, EL - 0.02, rows[2] + DH / 2.0)
label(CX + DW / 2.0, rows[2] + DH / 2.0 - 0.26, 1.2, "yes", align=PP_ALIGN.LEFT)
box(EL, rows[2] + DH / 2.0 - EH / 2.0, EW, EH,
    ["λ ← 2λ,  re-solve", "threshold too low to reward anything"],
    fill=PAPER, edge=LINE)
label(CX + 0.08, rows[2] + DH - 0.10, 1.0, "no", align=PP_ALIGN.LEFT)

# 4 diamond: converged
arrow(CX, rows[2] + DH, CX, rows[3] - 0.02)
box(L, rows[3], DW, DH, ["| objective |  ≤  tolerance ?"],
    shape=MSO_SHAPE.DIAMOND, size=10.5)
arrow(CX + DW / 2.0, rows[3] + DH / 2.0, EL - 0.02, rows[3] + DH / 2.0)
label(CX + DW / 2.0, rows[3] + DH / 2.0 - 0.26, 1.2, "yes", align=PP_ALIGN.LEFT)
box(EL, rows[3] + DH / 2.0 - EH / 2.0, EW, EH,
    ["STOP — break-even reached", "no surplus left to price out"],
    fill=PANEL, edge=LINE, bold=True)
label(CX + 0.08, rows[3] + DH - 0.24, 1.0, "no", align=PP_ALIGN.LEFT)

# 5 update
arrow(CX, rows[3] + DH, CX, rows[4] - 0.02)
box(L, rows[4], BW, 0.60,
    ["λₙₑₓₜ  =  D* / V*",
     "Dinkelbach: price at what was actually achieved"],
    fill=RGBColor(0xE9, 0xEE, 0xFB), edge=NAVY, color=NAVY, bold=True, width=1.5)

# 6 trial solve
arrow(CX, rows[4] + 0.60, CX, rows[5] - 0.02)
box(L, rows[5], BW, 0.52,
    ["Trial solve at λₙₑₓₜ", "recover the trial solution's V*"],
    fill=PAPER, edge=LINE)

# 7 diamond: degeneracy guard
arrow(CX, rows[5] + 0.52, CX, rows[6] - 0.02)
box(L, rows[6], DW, DH, ["trial V*  ≤  0 ?"], shape=MSO_SHAPE.DIAMOND,
    edge=WARN, color=WARN, size=11.5, width=1.6)
arrow(CX + DW / 2.0, rows[6] + DH / 2.0, EL - 0.02, rows[6] + DH / 2.0,
      color=WARN, width=1.4)
label(CX + DW / 2.0, rows[6] + DH / 2.0 - 0.26, 1.2, "yes", color=WARN,
      align=PP_ALIGN.LEFT)
box(EL, rows[6] + DH / 2.0 - 0.34, EW, 0.68,
    ["STOP — keep previous λ and solution",
     "guards against overshooting into “do nothing”"],
    fill=WARNBG, edge=WARN, color=WARN, bold=True, width=1.4)

# 8 adopt + cap, placed to the LEFT so the column does not run off the slide
label(L - 1.02, rows[6] + DH / 2.0 - 0.28, 1.0, "no", align=PP_ALIGN.RIGHT)
arrow(L, rows[6] + DH / 2.0, L - 0.92, rows[6] + DH / 2.0)
ADOPT_L, ADOPT_W = 0.70, 3.10
box(ADOPT_L, rows[6] + DH / 2.0 - 0.30, ADOPT_W, 0.60,
    ["Adopt  λ ← λₙₑₓₜ",
     "stop if the iteration cap is reached"],
    fill=PAPER, edge=LINE)

# loop back up the left margin to the solve step
LOOP_X = ADOPT_L + ADOPT_W / 2.0
elbow([(LOOP_X, rows[6] + DH / 2.0 - 0.30),
       (LOOP_X, rows[1] + 0.31),
       (L - 0.02, rows[1] + 0.31)])
label(LOOP_X + 0.12, rows[3] + 0.18, 2.6, "otherwise iterate again",
      align=PP_ALIGN.LEFT, italic=True)

out = os.path.join(W, "M4G-lambda-loop.pptx")
prs.save(out)
print("saved", out)
