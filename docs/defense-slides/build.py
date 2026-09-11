import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

W = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(W, "img")

NAVY  = RGBColor(0x1E, 0x27, 0x61)
ICE   = RGBColor(0xCA, 0xDC, 0xFC)
PAPER = RGBColor(0xF4, 0xF6, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY  = RGBColor(0x5B, 0x63, 0x7D)
GOLD  = RGBColor(0xC8, 0x8B, 0x2E)

SW, SH = 13.333, 7.5
HEAD  = "Cambria"
BODY  = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
BLANK = prs.slide_layouts[6]


def slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    return s


def box(s, x, y, w, h, text=None, size=16, color=NAVY, bold=False, font=BODY,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=6, italic=False):
    # Height is optional: box(s, x, y, w, text) is accepted and given a default
    # height, since a word-wrapped text box grows on its own anyway.
    if text is None and isinstance(h, (str, list)):
        text, h = h, 0.6
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        r = p.add_run()
        r.text = ln
        r.font.size, r.font.bold, r.font.italic = Pt(size), bold, italic
        r.font.color.rgb = color
        r.font.name = font
    return tb


def bullets(s, x, y, w, items, size=15, color=NAVY, gap=11, dot=GOLD, dotsize=0.075):
    """Bulleted list drawn with real dots, so no theme bullet glyphs leak in."""
    cy = y
    for it in items:
        d = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(cy + 0.085),
                               Inches(dotsize), Inches(dotsize))
        d.fill.solid(); d.fill.fore_color.rgb = dot
        d.line.fill.background(); d.shadow.inherit = False
        tb = box(s, x + 0.24, cy, w - 0.24, 0.4, it, size=size, color=color)
        tb.text_frame.word_wrap = True
        est = 1 + int(len(it) / max(1, int((w - 0.24) * 7.6 * (16.0 / size))))
        cy += est * (size / 72.0 * 1.32) + gap / 72.0
    return cy


def title(s, text, kicker=None):
    if kicker:
        box(s, 0.75, 0.44, 11.8, 0.3, kicker.upper(), size=11, color=GOLD,
            bold=True, font=BODY)
    box(s, 0.75, 0.76, 11.8, 0.75, text, size=30, color=NAVY, bold=True, font=HEAD)


def card(s, x, y, w, h, fill=PAPER, line=None):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                            Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(1)
    sh.shadow.inherit = False
    sh.adjustments[0] = 0.04
    if sh.has_text_frame:
        sh.text_frame.text = ""
    return sh


def formula(s, name, cx, cy, maxw, maxh):
    """Place a rendered formula centred on (cx, cy), scaled to fit the box."""
    p = os.path.join(IMG, name + ".png")
    iw, ih = Image.open(p).size
    ar = iw / float(ih)
    w, h = maxw, maxw / ar
    if h > maxh:
        h, w = maxh, maxh * ar
    s.shapes.add_picture(p, Inches(cx - w / 2.0), Inches(cy - h / 2.0),
                         Inches(w), Inches(h))
    return w, h


def stat(s, x, y, w, big, label, big_color=NAVY):
    box(s, x, y, w, 0.75, big, size=40, color=big_color, bold=True, font=HEAD,
        align=PP_ALIGN.CENTER)
    box(s, x, y + 0.72, w, 0.5, label, size=11.5, color=GREY,
        align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════ 1 · title
s = slide(NAVY)
box(s, 1.1, 2.25, 11.2, 0.4, "PROPOSED EXACT MODEL", size=13, color=ICE,
    bold=True, font=BODY)
box(s, 1.1, 2.75, 11.2, 1.5,
    "M4G — Online Joint OA–PS–TA Order-Splitting Model",
    size=40, color=WHITE, bold=True, font=HEAD)
box(s, 1.1, 4.45, 10.4, 1.3, [
    "Minimise distance per unit of order progress, re-solved every epoch.",
    "Two measured prices. Zero tuned constants."],
    size=17, color=ICE, space=8)
s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.1), Inches(4.22),
                   Inches(1.5), Inches(0.035)).fill.solid()
s.shapes[-1].fill.fore_color.rgb = GOLD
s.shapes[-1].line.fill.background(); s.shapes[-1].shadow.inherit = False
s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.1), Inches(1.35),
                   Inches(0.42), Inches(0.42)).fill.solid()
s.shapes[-1].fill.fore_color.rgb = GOLD
s.shapes[-1].line.fill.background(); s.shapes[-1].shadow.inherit = False
s.notes_slide.notes_text_frame.text = (
    "M4G decides order assignment, pod selection and task allocation in one MILP, "
    "once per decision epoch.")

# ══════════════════════════════════════════════════════════ 2 · what it decides
s = slide()
title(s, "One solve decides all three layers", "The decision")
card(s, 0.75, 1.85, 3.75, 1.85)
box(s, 1.05, 2.1, 3.15, 0.35, "OA", size=15, color=GOLD, bold=True, font=HEAD)
box(s, 1.05, 2.5, 3.15, 1.0, "Which orders occupy which station slots",
    size=14.5, color=NAVY)
card(s, 4.79, 1.85, 3.75, 1.85)
box(s, 5.09, 2.1, 3.15, 0.35, "PS", size=15, color=GOLD, bold=True, font=HEAD)
box(s, 5.09, 2.5, 3.15, 1.0, "Which pods are fetched, and to which station",
    size=14.5, color=NAVY)
card(s, 8.83, 1.85, 3.75, 1.85)
box(s, 9.13, 2.1, 3.15, 0.35, "TA", size=15, color=GOLD, bold=True, font=HEAD)
box(s, 9.13, 2.5, 3.15, 1.0, "Which robot is assigned to fetch each pod",
    size=14.5, color=NAVY)
formula(s, "sets", 6.67, 5.05, 11.0, 2.1)
box(s, 0.75, 6.45, 11.8,
    "Sequential OA → PS → TA cannot price a pod trip against the work it "
    "unlocks. M4G solves them jointly, so travel and progress are traded off inside "
    "one objective.", size=14, color=GREY)
s.notes_slide.notes_text_frame.text = (
    "Q(t) is filtered by stock: pod-SKU pairs with no inventory never become variables. "
    "That is what keeps the model tractable.")

# ══════════════════════════════════════════════════════════ 3 · objective
s = slide()
title(s, "Objective: distance per unit of progress", "Objective function")
card(s, 0.75, 1.8, 11.8, 1.6, ICE)
formula(s, "obj_ratio", 6.67, 2.6, 7.2, 1.05)
bullets(s, 0.9, 3.75, 11.5, [
    "D is the travel this decision creates, in metres — charged only for pods "
    "dispatched now; sunk trips are already paid for.",
    "V is the progress this decision is credited with, priced in metres by the two "
    "measured rates.",
    "The ratio is scale-free: it is the marginal cost of one unit of progress, so it "
    "stays comparable as the fleet, backlog and congestion change.",
    "Dinkelbach turns the fractional programme into a sequence of linear ones — "
    "no approximation, the fixed point is the exact minimiser.",
], size=15)
s.notes_slide.notes_text_frame.text = (
    "Say metres per line closed. Do not say cost function; the objective has no "
    "monetary units.")

# ══════════════════════════════════════════════════════════ 4 · D
s = slide()
title(s, "D — travel this decision creates", "Objective function · numerator")
card(s, 0.75, 1.8, 11.8, 1.75, PAPER)
formula(s, "obj_D", 6.67, 2.68, 10.6, 1.2)
box(s, 1.55, 3.72, 4.2, 0.35, "Pod → station leg", size=13, color=GOLD,
    bold=True, align=PP_ALIGN.CENTER)
box(s, 7.35, 3.72, 4.2, 0.35, "Robot → pod leg", size=13, color=GOLD,
    bold=True, align=PP_ALIGN.CENTER)
bullets(s, 0.9, 4.45, 11.5, [
    "Both sums range over P_a(t) only — pods already en route belong to P_b(t) "
    "and carry no charge, because that distance is sunk.",
    "This is what makes the model prefer to keep drawing from a pod that is already "
    "on its way: extra units from it are free.",
    "Distance is a validated proxy for mechanical energy: 80.65 ± 1.32 J/m within "
    "a fixed physical setting (cv 1.6%).",
], size=15)
s.notes_slide.notes_text_frame.text = (
    "The sunk/new split is structural, carried by the P_a subscript. No separate "
    "sunk-cost price is needed.")

# ══════════════════════════════════════════════════════════ 5 · V
s = slide()
title(s, "V — progress, converted to metres", "Objective function · denominator")
card(s, 0.75, 1.8, 11.8, 1.45, PAPER)
formula(s, "obj_V", 6.67, 2.5, 8.4, 1.0)
box(s, 0.75, 3.45, 11.8, [
    "The first term counts the lines the valuation layer expects to close, priced at "
    "λ (metres per line); the second counts the orders it expects to complete in "
    "full, priced at μ (metres per order).",
    "Both are therefore in metres, which is what makes D − V a meaningful "
    "difference."], size=15, color=NAVY, space=7)
card(s, 0.75, 4.75, 11.8, 1.35, ICE)
formula(s, "prices", 6.67, 5.42, 7.0, 0.95)
box(s, 0.75, 6.35, 11.8,
    "Both prices are running measurements of the system's own history — nothing "
    "is tuned. Completion is rewarded on top of the lines it contains, so the model "
    "has a reason to finish an order rather than close one line each across several.",
    size=14, color=GREY)
s.notes_slide.notes_text_frame.text = (
    "If asked: a completed order earns lambda per line AND mu, roughly double. That is "
    "deliberate, not double counting by accident.")

# ══════════════════════════════════════════════════════════ 6 · two layers
s = slide()
title(s, "Two layers: what is worth fetching vs. what fits", "Model structure")
card(s, 0.75, 1.8, 5.8, 2.5, PAPER)
box(s, 1.1, 2.05, 5.1, 0.4, "Valuation layer   q̂, ê, f̂",
    size=17, color=GOLD, bold=True, font=HEAD)
box(s, 1.1, 2.55, 5.1, 1.6, [
    "Not limited by station slots.",
    "Answers: is this pod worth fetching at all?",
    "A pod that can serve five lines is worth the trip even if only one slot is free "
    "now."], size=14, color=NAVY, space=5)
card(s, 6.79, 1.8, 5.76, 2.5, PAPER)
box(s, 7.14, 2.05, 5.1, 0.4, "Binding layer   q, e, f", size=17, color=NAVY,
    bold=True, font=HEAD)
box(s, 7.14, 2.55, 5.1, 1.6, [
    "Limited by station slots (B3).",
    "Answers: what can actually be committed this epoch?",
    "Only these draws have physical effect."], size=14, color=NAVY, space=5)
card(s, 0.75, 4.5, 11.8, 1.1, ICE)
formula(s, "layers", 6.67, 5.05, 6.2, 0.72)
box(s, 0.75, 5.85, 11.8, [
    "Separating them stops station congestion from suppressing a good pod choice — "
    "the classic failure of slot-limited greedy assignment.",
    "V is built entirely from the valuation layer, so a pod is priced by the work it "
    "makes possible, not only by what fits today."],
    size=14.5, color=NAVY, space=6)
s.notes_slide.notes_text_frame.text = (
    "The hat means expected, not delivered. Every V term is a hatted variable.")

# ══════════════════════════════════════════════════════════ 7 · lexicographic
s = slide()
title(s, "So why does anything actually get picked?", "Model structure · stage 2")
box(s, 0.75, 1.75, 11.8,
    "V contains only valuation variables — the binding layer earns nothing in the "
    "objective. Execution is forced by structure, not by a price.", size=15.5,
    color=NAVY)
card(s, 0.75, 2.5, 11.8, 1.15, PAPER)
formula(s, "b4", 6.67, 3.07, 8.6, 0.8)
box(s, 0.75, 3.75, 11.8,
    "B4 — no empty binding: claiming a slot requires at least one unit actually "
    "drawn.", size=14, color=GOLD, bold=True)
card(s, 0.75, 4.35, 11.8, 1.5, ICE)
formula(s, "lex", 6.67, 5.1, 8.6, 1.15)
bullets(s, 0.9, 6.05, 11.5, [
    "Minimising idle slots therefore maximises Σx, and every x needs a real draw "
    "— zero parameters, and it never enters the objective, so λ stays clean.",
    "Stage 2 cannot be infeasible: the stage-1 optimum is itself a witness.",
], size=14)
s.notes_slide.notes_text_frame.text = (
    "This replaces the discount factor beta used in the earlier version. Slot filling "
    "is a priority, not a price.")

# ══════════════════════════════════════════════════════════ 8 · Dinkelbach
s = slide()
title(s, "Dinkelbach is also the threshold controller", "Why a ratio, not a linear objective")
card(s, 0.75, 1.8, 11.8, 1.2, PAPER)
formula(s, "dinkel", 6.67, 2.4, 7.6, 0.85)
box(s, 0.75, 3.15, 11.8,
    "It tightens the acceptance threshold from “no worse than history” to "
    "“the best ratio currently reachable”. Measured by removing it:",
    size=15, color=NAVY)
for i, (x, big, lab) in enumerate([
        (0.75, "11.21 → 3.80", "λ: historical mean → current optimum"),
        (4.79, "110 → 519", "pod trips without Dinkelbach (4.7×)"),
        (8.83, "+38%", "total travel distance"),
]):
    card(s, x, 3.85, 3.75, 1.55, ICE)
    stat(s, x + 0.2, 4.05, 3.35, big, lab, big_color=GOLD if i == 0 else NAVY)
card(s, 0.75, 5.65, 11.8, 1.05, PAPER)
box(s, 1.1, 5.82, 11.1,
    "The denominator can never be negative — the null solution is feasible at "
    "objective 0, so:", size=13.5, color=GREY)
formula(s, "vstar", 6.67, 6.5, 6.6, 0.42)
s.notes_slide.notes_text_frame.text = (
    "The V* >= 0 line is the answer to 'can the denominator be negative'. It cannot.")

# ══════════════════════════════════════════════════════════ 9 · splitting
s = slide()
title(s, "What makes this a splitting model", "Contribution")
card(s, 0.75, 1.85, 11.8, 1.35, ICE)
formula(s, "split", 6.67, 2.52, 9.4, 0.95)
bullets(s, 0.9, 3.45, 11.5, [
    "The draw variable carries both a pod index p and a station index w, so one "
    "order's demand may be served by several pods, at several stations, across "
    "several epochs.",
    "M1G assigns a whole order to one station, all-or-nothing. Relaxing that to "
    "unit level is the modelling contribution.",
    "Splitting raises pile-on: a pod trip serves more lines, so the same items move "
    "for less travel.",
    "The cost is work-in-process — partially served orders wait for "
    "consolidation, which the model tracks and the experiments quantify.",
], size=15)
s.notes_slide.notes_text_frame.text = (
    "Compare items handled rather than orders completed: splitting creates WIP, which "
    "systematically understates completed orders over a finite horizon.")

# ══════════════════════════════════════════════════════════ 10 · variables
s = slide()
title(s, "Decision variables", "Formulation")
card(s, 0.75, 1.8, 11.8, 3.9, PAPER)
formula(s, "vars", 6.67, 3.75, 9.9, 3.4)
box(s, 0.75, 5.95, 11.8, [
    "A hat marks the valuation layer; the same symbol without one is the binding "
    "layer. Every term in V is hatted.",
    "The idle-slot variable exists only for stations that still have free "
    "capacity — a full station never enters the model at all."], size=14, color=GREY, space=6)
s.notes_slide.notes_text_frame.text = "Hatted = valuation layer; unhatted = binding layer."

# ══════════════════════════════════════════════════════════ 11 · constraints V
s = slide()
title(s, "Constraints — valuation layer", "Formulation")
formula(s, "cons_V", 6.67, 3.95, 11.6, 3.6)
box(s, 0.75, 6.15, 11.8, [
    "V1 ties every draw to a dispatched pod and its stock · V2 caps a line at its "
    "residual demand · V2a discounts inbound replenishment already on its way",
    "V3 credits a line only when it is drawn in full · V4 credits an order only "
    "when every one of its lines is closed"], size=13.5, color=GREY, space=5)

# ══════════════════════════════════════════════════════════ 12 · constraints R
s = slide()
title(s, "Constraints — resources", "Formulation")
formula(s, "cons_R", 6.67, 3.85, 11.0, 3.4)
box(s, 0.75, 6.0, 11.8, [
    "R1 one station per pod · R2 a dispatch needs a carrier · R3 one pod per "
    "robot · R4 one robot per pod",
    "R5 pins pods already en route — their assignment is a physical fact, not a "
    "decision"], size=13.5, color=GREY, space=5)

# ══════════════════════════════════════════════════════════ 13 · constraints B
s = slide()
title(s, "Constraints — binding layer", "Formulation")
formula(s, "cons_B", 6.67, 4.05, 11.6, 4.0)
box(s, 0.75, 6.35, 11.8,
    "B1/B6 make the binding layer a subset of the valuation layer · B2 and B4 "
    "together give “drawing ⇔ occupying a slot” · B3s defines the "
    "idle-slot variable stage 2 minimises", size=13.5, color=GREY)

# ══════════════════════════════════════════════════════════ 14 · closing
s = slide(NAVY)
box(s, 0.9, 0.85, 11.5, 0.4, "SUMMARY", size=13, color=GOLD, bold=True)
box(s, 0.9, 1.25, 11.5, 0.9, "Two measured prices, zero tuned constants",
    size=32, color=WHITE, bold=True, font=HEAD)
rows = [
    ("λ", "metres per line closed", "measured from cumulative history"),
    ("μ", "metres per order completed", "measured from cumulative history"),
    ("ρ", "pod-tier sunk-cost price", "removed — already carried by the pod-set subscript"),
    ("ε", "unit-level tie-break", "removed — the only hand-set constant"),
    ("β", "realisation-rate discount", "replaced by the stage-2 tie-break"),
]
y = 2.5
for i, (sym, what, status) in enumerate(rows):
    kept = i < 2
    card(s, 0.9, y, 11.5, 0.72,
         RGBColor(0x2B, 0x35, 0x76) if kept else RGBColor(0x25, 0x2C, 0x62))
    box(s, 1.25, y + 0.14, 0.7, 0.45, sym, size=21,
        color=GOLD if kept else RGBColor(0x6E, 0x79, 0x9E), bold=True, font=HEAD)
    box(s, 2.15, y + 0.2, 4.3, 0.4, what, size=14.5,
        color=WHITE if kept else RGBColor(0x8F, 0x99, 0xBA))
    box(s, 6.7, y + 0.2, 5.4, 0.4, status, size=14.5,
        color=ICE if kept else RGBColor(0x6E, 0x79, 0x9E),
        italic=not kept)
    y += 0.82
box(s, 0.9, 6.75, 11.5,
    "Every remaining price is a measurement of the system's own behaviour, so the "
    "model has nothing left to tune.", size=14.5, color=ICE)

out = os.path.join(W, "M4G-model.pptx")
prs.save(out)
print("saved", out, len(prs.slides.__iter__.__self__._sldIdLst), "slides")
